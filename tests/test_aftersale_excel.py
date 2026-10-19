# -*- coding: utf-8 -*-
"""database.aftersale_excel 回归（A3 excel 切片）。"""
import os
import sys

import pytest

pytest.importorskip("openpyxl")

from openpyxl import Workbook

from database import aftersale_db as adb
from database import aftersale_excel as ax

TUE = {"type": "tue", "start": "", "span": 7}


def test_safe_sheet_title():
    assert ax._safe_sheet_title("硬件:问题/类型") == "硬件-问题-类型"
    assert ax._safe_sheet_title("a" * 40) == "a" * 31
    assert ax._safe_sheet_title("") == "未分类"


def test_write_export_sheet_headers_cycle_and_freeze(monkeypatch):
    monkeypatch.setattr(adb, "load_cycle_mode", lambda: dict(TUE))
    wb = Workbook()
    ws = wb.active
    rows = [{"issue_type": "硬件问题", "room_name": "甲",
             "occurred_at": "2026-08-19 10:00:00", "created_at": "2026-08-19"}]
    ax._write_export_sheet(ws, rows)
    headers = [ws.cell(row=1, column=c).value for c in range(1, 15)]
    assert headers[0] == "类型"
    assert headers[-1] == "周期"
    # 周期列按发生时间动态计算，与列表口径一致
    cycle_col = headers.index("周期") + 1
    assert ws.cell(row=2, column=cycle_col).value == "08/18 - 08/24"
    assert ws.freeze_panes == "A2"


def test_parse_excel_rows_fills_type_and_skips_blank(monkeypatch, tmp_path):
    monkeypatch.setattr(adb, "load_cycle_mode", lambda: dict(TUE))
    wb = Workbook()
    ws = wb.active
    ws.append(["类型", "球房", "桌号", "地区", "问题", "是否解决"])
    ws.append(["硬件问题", "甲", "T1", "广东", "不亮", "是"])
    ws.append(["", "甲", "T2", "广东", "异响", ""])   # 类型向下填充，解决空→否
    ws.append(["", "", "", "", "", ""])               # 全空跳过
    path = tmp_path / "in.xlsx"
    wb.save(path)

    headers, data = ax.parse_excel_rows(str(path))
    assert "类型" in headers and "问题" in headers
    assert len(data) == 2
    assert data[0]["issue_type"] == "硬件问题"
    assert data[1]["issue_type"] == "硬件问题"
    assert data[1]["resolved"] == "否"
    assert data[0]["creator"] == "Excel导入"
    assert data[0]["cycle_start"] == adb.cycle_start_of(
        __import__("datetime").datetime.now())
