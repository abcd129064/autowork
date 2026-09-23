# -*- coding: utf-8 -*-
"""售后排行页（球房/球桌两级排行，与 Web 端 /aftersale/rank 同功能）

RankPage：售后面板 FluentWindow 第 4 个导航子页。信息架构对齐 Web 端：
① 筛选栏（时间范围/账期/地区/三个是否/关键词）② 概览 4 卡（覆盖球房/
覆盖球桌/记录总数/未解决）③ 球房/球桌分段 + TOP 深度 + 排序 + 面包屑
下钻 ④ 左 QPainter 自绘横向条形图 + 右表格（排名徽章 + 指标列 + 操作列
文字链接「明细/下钻」）。

数据经 ``aftersale_db.query_rank`` 一次取回（异步 Worker + _query_seq
代数守卫，与 RecordsPage 同范式）；图表 QPainter 自绘零新依赖（与
AfterSaleStatsDialog 同范式），深浅主题自适应。下钻为服务端口径
（query_rank room_name 参数），面包屑「全部球房」一键返回全局榜。
明细跳转发 jump_to_records(keyword) 信号，由宿主窗口切换记录页并
set_keyword 预筛选（球房级=球房名，球桌级=桌号）。
"""

from datetime import date, timedelta

from PySide6.QtCore import Qt, QRect, QRectF, QSize, QTimer, Signal
from PySide6.QtGui import QColor, QFont, QFontMetrics, QPainter
from PySide6.QtWidgets import (QAbstractItemView, QFileDialog, QFrame,
    QHBoxLayout, QHeaderView, QLabel, QScrollArea, QSizePolicy,
    QTableWidgetItem, QVBoxLayout, QWidget)

from qfluentwidgets import (BodyLabel, CaptionLabel, CardWidget, ComboBox,
    FluentIcon, FlowLayout, PrimaryPushButton, PushButton, SearchLineEdit,
    SegmentedWidget, TableWidget, ZhDatePicker,
    isDarkTheme)

from core.design_tokens import SEMANTIC
from core.flow_widgets import FlowToolbarScrollArea
from core.ops_link_delegate import LINKS_ROLE, install_ops_links
from core.perf import apply_table_smooth_mode
from core.theme_qss import current_accent_hex
from core.utils import show_info_bar
from database import aftersale_db
from workers.aftersale_worker import AftersaleDBWorker

from windows.aftersale.common import (  # noqa: F401
    _FIXED_ROW_HEIGHT, YesNoSegment,
)

# 主题取色辅助（与 stats_dialog 同口径：浅/深主题各一套）
def _text_color() -> str:
    return "#6B7280" if not isDarkTheme() else "#8892a2"


def _muted_color() -> str:
    return "#9CA3AF" if not isDarkTheme() else "#555f6b"


# 排名前三徽章色（金/银/铜，深浅主题通用的中明度色）
_MEDAL_COLORS = ("#d4af37", "#9aa3ad", "#c08457")

# 表格列（列宽 0 = Stretch 拉伸列）
_RANK_COLUMNS = (
    ("rank", "排名", 46),
    ("name", "名称", 0),
    ("total", "总量", 60),
    ("share", "占比", 56),
    ("unresolved", "未解决", 64),
    ("our_problem", "我方问题", 76),
    ("last_occurred", "最近发生", 92),
    ("ops", "操作", 108),
)
_COL_NAME = 1
_COL_OPS = len(_RANK_COLUMNS) - 1


def _fetch_rank_options():
    """排行下拉选项（周期 + 地区），供 AftersaleDBWorker 后台线程调用"""
    try:
        regions = list(aftersale_db.get_field_candidates().get("regions") or [])
    except Exception:
        regions = []
    return aftersale_db.get_cycle_options(), regions


# ==================== 排行横向条形图（QPainter 自绘） ====================

class _RankHBarChart(QWidget):
    """排行横向条形图：名称 | 条形 | 数量 · 占比（TOP N 行，降序输入）

    行可点击：rowClicked 发射该行序号（与页面 _rows 同序），由页面分发
    （球房级行点击=下钻该球房，球桌级行点击=跳转明细）。悬停行淡色高亮。
    """

    ROW_H = 28
    LABEL_W = 150
    VAL_W = 96
    rowClicked = Signal(int)

    def __init__(self, parent=None):
        super().__init__(parent)
        self._rows = []
        self._bar_color = None    # None=主题强调色（球房榜）
        self._hover = -1
        self.setMouseTracking(True)
        self.setCursor(Qt.CursorShape.PointingHandCursor)
        self.setSizePolicy(QSizePolicy.Policy.Expanding,
                           QSizePolicy.Policy.Expanding)

    def set_rows(self, rows, bar_color=None):
        self._rows = list(rows or [])
        self._bar_color = bar_color
        self._hover = -1
        self.setMinimumHeight(
            len(self._rows) * self.ROW_H + 10 if self._rows else 64)
        self.update()

    def sizeHint(self):
        return QSize(320, max(64, len(self._rows) * self.ROW_H + 10))

    def mouseMoveEvent(self, e):
        idx = -1
        if self._rows:
            idx = int(e.position().y()) // self.ROW_H
            if not (0 <= idx < len(self._rows)):
                idx = -1
        if idx != self._hover:
            self._hover = idx
            self.update()
        super().mouseMoveEvent(e)

    def leaveEvent(self, e):
        if self._hover != -1:
            self._hover = -1
            self.update()
        super().leaveEvent(e)

    def mousePressEvent(self, e):
        if e.button() == Qt.MouseButton.LeftButton and self._rows:
            idx = int(e.position().y()) // self.ROW_H
            if 0 <= idx < len(self._rows):
                self.rowClicked.emit(idx)
        super().mousePressEvent(e)

    def paintEvent(self, _e):
        p = QPainter(self)
        p.setRenderHint(QPainter.RenderHint.Antialiasing)
        if not self._rows:
            p.setPen(QColor(_text_color()))
            p.drawText(self.rect(), Qt.AlignmentFlag.AlignCenter,
                       "当前筛选下暂无排行数据")
            return
        fm = QFontMetrics(p.font())
        max_v = max(int(r.get("total") or 0) for r in self._rows) or 1
        base = QColor(self._bar_color or current_accent_hex())
        hover_c = QColor(base)
        hover_c.setAlpha(28)
        bar_x = self.LABEL_W + 8
        bar_w = max(20.0, self.width() - bar_x - self.VAL_W - 8)
        for i, r in enumerate(self._rows):
            y = 5 + i * self.ROW_H
            if i == self._hover:
                p.setPen(Qt.PenStyle.NoPen)
                p.setBrush(hover_c)
                p.drawRoundedRect(
                    QRect(2, y - 3, self.width() - 4, self.ROW_H - 2), 4, 4)
            n = int(r.get("total") or 0)
            name = fm.elidedText(
                str(r.get("name") or ""), Qt.TextElideMode.ElideRight,
                self.LABEL_W - 6)
            p.setPen(QColor(_text_color()))
            p.drawText(QRect(0, y, self.LABEL_W, 18),
                       Qt.AlignmentFlag.AlignRight
                       | Qt.AlignmentFlag.AlignVCenter, name)
            bw = max(3.0, bar_w * n / max_v)
            p.setPen(Qt.PenStyle.NoPen)
            p.setBrush(base)
            p.drawRoundedRect(QRectF(bar_x, y + 2, bw, 13), 4, 4)
            p.setPen(QColor(_muted_color()))
            p.drawText(
                QRectF(bar_x + bar_w + 8, y, self.VAL_W, 18),
                Qt.AlignmentFlag.AlignLeft | Qt.AlignmentFlag.AlignVCenter,
                f"{n} · {int(r.get('share') or 0)}%")


# ==================== 售后排行页 ====================

class RankPage(QWidget):
    """售后排行页：筛选栏 + 概览 4 卡 + 球房/球桌排行（左图右表 + 下钻）

    明细跳转经 jump_to_records(keyword) 由宿主窗口接线（切换记录页并
    关键词预筛选）；下钻为服务端口径（query_rank room_name），面包屑
    返回全局榜。
    """

    jump_to_records = Signal(str)   # 明细跳转：携带预筛选关键词

    _PERF_PANEL_KEY = "aftersale"
    _LIMITS = (10, 20, 50, 200)
    _SORTS = (("total", "按总量"), ("unresolved", "按未解决"),
              ("our_problem", "按我方问题"), ("last_occurred", "按最近发生"))

    def __init__(self, parent=None):
        super().__init__(parent)
        self._worker = None
        self._options_worker = None
        self._query_seq = 0
        self._rows = []
        self._summary = {}
        self._level = "room"      # room | table
        self._drill = None        # 下钻球房名（None=全局榜）
        self._options_loaded = False
        self._init_ui()
        self._search_timer = QTimer(self)
        self._search_timer.setInterval(300)
        self._search_timer.setSingleShot(True)
        self._search_timer.timeout.connect(self._on_filter_changed)

    # ---------- UI 构造 ----------

    def _init_ui(self):
        root = QVBoxLayout(self)
        root.setContentsMargins(20, 14, 20, 12)
        root.setSpacing(8)

        # --- 筛选工具栏（与记录页同范式：FlowLayout 换行 + 滚动容器锁高） ---
        toolbar_scroll = FlowToolbarScrollArea(self)
        toolbar_scroll.setWidgetResizable(True)
        toolbar_scroll.setHorizontalScrollBarPolicy(
            Qt.ScrollBarPolicy.ScrollBarAlwaysOff)
        toolbar_scroll.setVerticalScrollBarPolicy(
            Qt.ScrollBarPolicy.ScrollBarAsNeeded)
        toolbar_scroll.setMaximumHeight(132)
        toolbar_widget = QWidget(self)
        toolbar_scroll.setWidget(toolbar_widget)
        toolbar = FlowLayout(toolbar_widget)
        toolbar.setHorizontalSpacing(6)
        toolbar.setVerticalSpacing(6)
        toolbar.setContentsMargins(2, 4, 2, 4)

        toolbar.addWidget(CaptionLabel("时间范围：", self))
        self._range_seg = SegmentedWidget(self)
        for key, text in (("d30", "近 30 天"), ("d90", "近 90 天"),
                          ("all", "全部"), ("cycle", "账期"),
                          ("custom", "自定义")):
            self._range_seg.addItem(key, text)
        self._range_seg.setCurrentItem("d90")
        self._range_seg.currentItemChanged.connect(self._on_range_changed)
        toolbar.addWidget(self._range_seg)

        # 账期下拉（仅「账期」可见）与自定义起止（仅「自定义」可见）
        self._cycle_combo = ComboBox(self)
        self._cycle_combo.setFixedWidth(170)
        self._cycle_combo.currentIndexChanged.connect(
            lambda _i: self._on_filter_changed())
        toolbar.addWidget(self._cycle_combo)
        self._start_picker = ZhDatePicker(self)
        self._start_picker.setFixedWidth(140)
        self._start_picker.setDate(self._qdate_today())
        self._end_picker = ZhDatePicker(self)
        self._end_picker.setFixedWidth(140)
        self._end_picker.setDate(self._qdate_today())
        try:
            self._start_picker.dateChanged.connect(
                lambda _q: self._on_filter_changed())
            self._end_picker.dateChanged.connect(
                lambda _q: self._on_filter_changed())
        except Exception:
            pass
        toolbar.addWidget(self._start_picker)
        self._lbl_to = CaptionLabel("至", self)
        toolbar.addWidget(self._lbl_to)
        toolbar.addWidget(self._end_picker)
        self._cycle_combo.hide()
        self._start_picker.hide()
        self._end_picker.hide()
        self._lbl_to.hide()

        toolbar.addWidget(CaptionLabel("地区：", self))
        self._region_combo = ComboBox(self)
        self._region_combo.setFixedWidth(120)
        self._region_combo.addItem("全部地区", userData="")
        self._region_combo.currentIndexChanged.connect(
            lambda _i: self._on_filter_changed())
        toolbar.addWidget(self._region_combo)

        self._resolved_seg = YesNoSegment("", self, include_all=True)
        self._resolved_seg.currentItemChanged.connect(
            lambda _k: self._on_filter_changed())
        toolbar.addWidget(CaptionLabel("是否解决：", self))
        toolbar.addWidget(self._resolved_seg)
        self._our_problem_seg = YesNoSegment("", self, include_all=True)
        self._our_problem_seg.setToolTip("是否是我们的问题")
        self._our_problem_seg.currentItemChanged.connect(
            lambda _k: self._on_filter_changed())
        toolbar.addWidget(CaptionLabel("我方问题：", self))
        toolbar.addWidget(self._our_problem_seg)
        self._initiative_seg = YesNoSegment("", self, include_all=True)
        self._initiative_seg.setToolTip("是否我们主动发起")
        self._initiative_seg.currentItemChanged.connect(
            lambda _k: self._on_filter_changed())
        toolbar.addWidget(CaptionLabel("主动发起：", self))
        toolbar.addWidget(self._initiative_seg)

        self._search_edit = SearchLineEdit(self)
        self._search_edit.setPlaceholderText("球房 / 桌号")
        self._search_edit.setFixedWidth(170)
        self._search_edit.textChanged.connect(
            lambda _t: self._search_timer.start())
        toolbar.addWidget(self._search_edit)

        self._btn_query = PrimaryPushButton(FluentIcon.SEARCH, "查询", self)
        self._btn_query.setToolTip("按当前筛选重新查询排行")
        self._btn_query.clicked.connect(self._load)
        toolbar.addWidget(self._btn_query)
        self._btn_reset = PushButton(FluentIcon.SYNC, "重置", self)
        self._btn_reset.setToolTip("恢复默认筛选（近 90 天 · 全部分类）")
        self._btn_reset.clicked.connect(self._on_reset)
        toolbar.addWidget(self._btn_reset)
        self._btn_export = PushButton(FluentIcon.DOWNLOAD, "导出 xlsx", self)
        self._btn_export.setToolTip("导出当前排行（含指标列）为 xlsx")
        self._btn_export.clicked.connect(self._on_export)
        toolbar.addWidget(self._btn_export)
        root.addWidget(toolbar_scroll)

        # --- 概览 4 卡（覆盖球房/覆盖球桌/记录总数/未解决） ---
        cards_row = QHBoxLayout()
        cards_row.setSpacing(10)
        self._card_rooms = self._make_card(cards_row)
        self._card_tables = self._make_card(cards_row)
        self._card_total = self._make_card(cards_row)
        self._card_unresolved = self._make_card(cards_row)
        root.addLayout(cards_row)

        # --- 排行卡（分段 + TOP/排序 + 面包屑 + 左图右表） ---
        card = CardWidget(self)
        card_lay = QVBoxLayout(card)
        card_lay.setContentsMargins(14, 10, 14, 12)
        card_lay.setSpacing(8)

        head = QHBoxLayout()
        head.setSpacing(8)
        self._lbl_rank_title = BodyLabel("球房售后排行", card)
        head.addWidget(self._lbl_rank_title)
        self._level_seg = SegmentedWidget(card)
        self._level_seg.addItem("room", "球房")
        self._level_seg.addItem("table", "球桌")
        self._level_seg.setCurrentItem("room")
        self._level_seg.currentItemChanged.connect(self._on_level_changed)
        head.addWidget(self._level_seg)
        # 面包屑（下钻时出现）：返回按钮 + 当前球房名
        self._btn_back = PushButton("‹ 全部球房", card)
        self._btn_back.setToolTip("返回全局球房排行")
        self._btn_back.clicked.connect(self._on_drill_back)
        self._btn_back.hide()
        head.addWidget(self._btn_back)
        self._lbl_drill = CaptionLabel("", card)
        self._lbl_drill.setStyleSheet(
            f"color: {current_accent_hex()}; background: transparent;")
        self._lbl_drill.hide()
        head.addWidget(self._lbl_drill)
        head.addStretch(1)
        head.addWidget(CaptionLabel("TOP", card))
        self._limit_combo = ComboBox(card)
        for n in self._LIMITS:
            self._limit_combo.addItem(str(n), userData=n)
        self._limit_combo.setFixedWidth(78)
        self._limit_combo.currentIndexChanged.connect(
            lambda _i: self._on_filter_changed())
        head.addWidget(self._limit_combo)
        head.addWidget(CaptionLabel("排序", card))
        self._sort_combo = ComboBox(card)
        for key, text in self._SORTS:
            self._sort_combo.addItem(text, userData=key)
        self._sort_combo.setFixedWidth(116)
        self._sort_combo.currentIndexChanged.connect(
            lambda _i: self._on_filter_changed())
        head.addWidget(self._sort_combo)
        card_lay.addLayout(head)

        body = QHBoxLayout()
        body.setSpacing(12)
        self._chart_scroll = QScrollArea(card)
        self._chart_scroll.setWidgetResizable(True)
        self._chart_scroll.setFrameShape(QFrame.Shape.NoFrame)
        self._chart_scroll.setHorizontalScrollBarPolicy(
            Qt.ScrollBarPolicy.ScrollBarAlwaysOff)
        self._chart_scroll.setVerticalScrollBarPolicy(
            Qt.ScrollBarPolicy.ScrollBarAsNeeded)
        self._chart = _RankHBarChart(self._chart_scroll)
        self._chart_scroll.setWidget(self._chart)
        self._chart.rowClicked.connect(self._on_chart_row)
        body.addWidget(self._chart_scroll, 2)

        self._table = TableWidget(card)
        install_ops_links(self._table, self._on_ops_link)
        apply_table_smooth_mode(self._table, panel=self._PERF_PANEL_KEY)
        self._table.setColumnCount(len(_RANK_COLUMNS))
        self._table.setHorizontalHeaderLabels(
            [c[1] for c in _RANK_COLUMNS])
        self._table.setSelectionMode(
            QAbstractItemView.SelectionMode.NoSelection)
        self._table.setEditTriggers(QAbstractItemView.EditTrigger.NoEditTriggers)
        self._table.setWordWrap(False)
        self._table.verticalHeader().setVisible(False)
        self._table.verticalHeader().setDefaultSectionSize(_FIXED_ROW_HEIGHT)
        header = self._table.horizontalHeader()
        header.setStretchLastSection(False)
        for i, (_k, _h, w) in enumerate(_RANK_COLUMNS):
            if w == 0:
                header.setSectionResizeMode(
                    i, QHeaderView.ResizeMode.Stretch)
            else:
                header.setSectionResizeMode(i, QHeaderView.ResizeMode.Fixed)
                self._table.setColumnWidth(i, w)
        body.addWidget(self._table, 3)
        card_lay.addLayout(body, 1)
        root.addWidget(card, 1)

    @staticmethod
    def _qdate_today():
        from PySide6.QtCore import QDate
        return QDate.currentDate()

    def _make_card(self, layout) -> tuple:
        """单张概览卡：标签 + 大数字 + 辅助说明（与记录页指标卡同款）"""
        card = CardWidget(self)
        lay = QVBoxLayout(card)
        lay.setContentsMargins(14, 10, 14, 10)
        lay.setSpacing(2)
        lbl = CaptionLabel("", card)
        num = QLabel("0", card)
        num.setStyleSheet(
            "font-size: 24px; font-weight: 500; background: transparent;")
        sub = CaptionLabel("", card)
        lay.addWidget(lbl)
        lay.addWidget(num)
        lay.addWidget(sub)
        layout.addWidget(card, 1)
        return (card, lbl, num, sub)

    # ---------- 筛选条件 ----------

    def _current_filters(self) -> dict:
        """汇集当前筛选（时间范围→start/end 或 cycle_start，分类三态，
        地区与关键词；下钻球房由 _drill 单独传递）"""
        rng = self._range_seg.currentRouteKey()
        start = end = ""
        cycle_start = ""
        if rng == "d30":
            start = (date.today() - timedelta(days=29)).isoformat()
            end = date.today().isoformat()
        elif rng == "d90":
            start = (date.today() - timedelta(days=89)).isoformat()
            end = date.today().isoformat()
        elif rng == "cycle":
            cycle_start = str(self._cycle_combo.currentData() or "")
        elif rng == "custom":
            start = self._start_picker.date.toString("yyyy-MM-dd")
            end = self._end_picker.date.toString("yyyy-MM-dd")
            if start and end and start > end:
                start, end = end, start
        return {
            "start": start, "end": end, "cycle_start": cycle_start,
            "resolved": self._resolved_seg.value(),
            "is_initiative": self._initiative_seg.value(),
            "is_our_problem": self._our_problem_seg.value(),
            "region": str(self._region_combo.currentData() or ""),
            "keyword": self._search_edit.text().strip(),
        }

    def _on_range_changed(self, _key):
        """时间范围切换：联动账期/自定义控件显隐后重查"""
        rng = self._range_seg.currentRouteKey()
        self._cycle_combo.setVisible(rng == "cycle")
        for w in (self._start_picker, self._end_picker, self._lbl_to):
            w.setVisible(rng == "custom")
        self._on_filter_changed()

    def _on_filter_changed(self):
        self._load()

    def _on_reset(self):
        """重置：恢复默认筛选（近 90 天/全部地区/三态全部/清关键词），
        blockSignals 后统一重查一次，避免逐控件触发多次查询"""
        for seg in (self._resolved_seg, self._our_problem_seg,
                    self._initiative_seg):
            seg.blockSignals(True)
            seg.setValue("")
            seg.blockSignals(False)
        self._range_seg.blockSignals(True)
        self._range_seg.setCurrentItem("d90")
        self._range_seg.blockSignals(False)
        for combo in (self._cycle_combo, self._region_combo,
                      self._limit_combo, self._sort_combo):
            combo.blockSignals(True)
            combo.setCurrentIndex(0)
            combo.blockSignals(False)
        self._search_edit.blockSignals(True)
        self._search_edit.clear()
        self._search_edit.blockSignals(False)
        self._cycle_combo.hide()
        for w in (self._start_picker, self._end_picker, self._lbl_to):
            w.hide()
        self._drill = None
        self._level_seg.blockSignals(True)
        self._level_seg.setCurrentItem("room")
        self._level_seg.blockSignals(False)
        self._level = "room"
        self._load()

    # ---------- 选项与数据加载 ----------

    def showEvent(self, event):
        super().showEvent(event)
        if not self._options_loaded:
            self._load_options()

    def reload_options(self):
        """周期配置保存后由宿主窗口调用：下拉选项置旧（可见时立即重拉）"""
        self._options_loaded = False
        if self.isVisible():
            self._load_options()

    def _load_options(self):
        if (self._options_worker is not None
                and self._options_worker.isRunning()):
            return
        self._options_worker = AftersaleDBWorker(_fetch_rank_options)
        self._options_worker.result_ready.connect(self._on_options_loaded)
        self._options_worker.error.connect(lambda _m: self._load())
        self._options_worker.start()

    def _on_options_loaded(self, opts):
        self._options_loaded = True
        cycles, regions = opts or ([], [])
        self._cycle_combo.blockSignals(True)
        self._cycle_combo.clear()
        self._cycle_combo.addItem("全部周期", userData="")
        for cs in cycles or []:
            self._cycle_combo.addItem(
                aftersale_db.cycle_label(cs), userData=cs)
        self._cycle_combo.blockSignals(False)
        self._region_combo.blockSignals(True)
        self._region_combo.clear()
        self._region_combo.addItem("全部地区", userData="")
        for rg in regions or []:
            self._region_combo.addItem(str(rg), userData=str(rg))
        self._region_combo.blockSignals(False)
        self._load()

    def _load(self):
        """按当前筛选异步查询排行（Worker + seq 代数守卫，同记录页范式）"""
        if self._worker and self._worker.isRunning():
            self._worker.requestInterruption()
        self._query_seq += 1
        seq = self._query_seq
        f = self._current_filters()
        self._worker = AftersaleDBWorker(
            aftersale_db.query_rank,
            self._level, int(self._limit_combo.currentData() or 10),
            str(self._sort_combo.currentData() or "total"),
            start=f["start"], end=f["end"], cycle_start=f["cycle_start"],
            resolved=f["resolved"], is_initiative=f["is_initiative"],
            is_our_problem=f["is_our_problem"], region=f["region"],
            room_name=self._drill or "", keyword=f["keyword"])
        self._worker.result_ready.connect(
            lambda result, _s=seq: self._on_loaded(result, _s))
        self._worker.error.connect(
            lambda msg, _s=seq: self._on_error(msg, _s))
        self._worker.start()

    def _on_loaded(self, result, seq=None):
        if seq is not None and seq != self._query_seq:
            return  # 过期查询结果（已被更新的筛选取代），丢弃
        self._rows = list((result or {}).get("rows") or [])
        self._summary = dict((result or {}).get("summary") or {})
        self._update_summary_cards()
        self._update_rank_title()
        bar = (SEMANTIC["warning"]
               if (self._drill or self._level == "table") else None)
        self._chart.set_rows(self._rows, bar_color=bar)
        self._populate_table()

    def _on_error(self, msg, seq=None):
        if seq is not None and seq != self._query_seq:
            return
        self._chart.set_rows([])
        self._table.setRowCount(0)
        self._lbl_rank_title.setText(f"售后排行（查询失败：{msg}）")

    # ---------- 渲染 ----------

    def _update_summary_cards(self):
        s = self._summary
        total = int(s.get("total") or 0)
        unresolved = int(s.get("unresolved") or 0)
        base = "font-size: 24px; font-weight: 500; background: transparent;"

        _c, lbl, num, sub = self._card_rooms
        lbl.setText("覆盖球房")
        num.setText(str(int(s.get("rooms") or 0)))
        num.setStyleSheet(base)
        sub.setText("去重非空球房数")

        _c, lbl, num, sub = self._card_tables
        lbl.setText("覆盖球桌")
        num.setText(str(int(s.get("tables") or 0)))
        num.setStyleSheet(base)
        sub.setText("去重球房·桌号对")

        _c, lbl, num, sub = self._card_total
        lbl.setText("记录总数")
        num.setText(str(total))
        num.setStyleSheet(base)
        sub.setText("当前筛选口径")

        _c, lbl, num, sub = self._card_unresolved
        lbl.setText("未解决")
        num.setText(str(unresolved))
        num.setStyleSheet(base + (
            f" color: {SEMANTIC['danger']};" if unresolved else ""))
        sub.setText("待处理积压" if unresolved else "全部已处理")

    def _update_rank_title(self):
        if self._drill:
            self._lbl_rank_title.setText("球桌售后排行")
            self._btn_back.show()
            self._lbl_drill.setText(self._drill)
            self._lbl_drill.show()
        else:
            self._lbl_rank_title.setText(
                "球房售后排行" if self._level == "room"
                else "球桌售后排行")
            self._btn_back.hide()
            self._lbl_drill.hide()

    def _populate_table(self):
        """排行行 → 表格：前三金银铜徽章 + 指标列语义色 + 操作列链接"""
        rows = self._rows
        table = self._table
        table.setUpdatesEnabled(False)
        table.blockSignals(True)
        try:
            table.setRowCount(len(rows))
            bold = QFont()
            bold.setBold(True)
            center = Qt.AlignmentFlag.AlignCenter
            for r, rec in enumerate(rows):
                rank = int(rec.get("rank") or (r + 1))
                it = QTableWidgetItem(str(rank))
                it.setTextAlignment(center)
                if rank <= 3:
                    it.setForeground(QColor(_MEDAL_COLORS[rank - 1]))
                    it.setFont(bold)
                table.setItem(r, 0, it)

                it = QTableWidgetItem(str(rec.get("name") or ""))
                if rec.get("table_no"):
                    it.setToolTip(f"球房: {rec.get('room_name') or ''}\n"
                                  f"桌号: {rec.get('table_no') or ''}")
                table.setItem(r, _COL_NAME, it)

                it = QTableWidgetItem(str(int(rec.get("total") or 0)))
                it.setTextAlignment(center)
                it.setFont(bold)
                table.setItem(r, 2, it)

                it = QTableWidgetItem(f"{int(rec.get('share') or 0)}%")
                it.setTextAlignment(center)
                table.setItem(r, 3, it)

                un = int(rec.get("unresolved") or 0)
                it = QTableWidgetItem(str(un))
                it.setTextAlignment(center)
                if un:
                    it.setForeground(QColor(SEMANTIC["danger"]))
                table.setItem(r, 4, it)

                our = int(rec.get("our_problem") or 0)
                it = QTableWidgetItem(str(our))
                it.setTextAlignment(center)
                if our:
                    it.setForeground(QColor(SEMANTIC["warning"]))
                table.setItem(r, 5, it)

                it = QTableWidgetItem(str(rec.get("last_occurred") or "—"))
                it.setForeground(QColor(_muted_color()))
                table.setItem(r, 6, it)

                # 操作列：链接清单按层级取（球房全局榜 2 段，其余 1 段）
                it = QTableWidgetItem("")
                links = [("明细", "primary", "detail")]
                if self._level == "room" and not self._drill:
                    links.append(("下钻", "primary", "drill"))
                it.setData(LINKS_ROLE, tuple(links))
                table.setItem(r, _COL_OPS, it)
        finally:
            table.setUpdatesEnabled(True)
            table.blockSignals(False)

    # ---------- 交互 ----------

    def _on_level_changed(self, key):
        """球房/球桌分段切换：清下钻并重查"""
        self._level = "table" if key == "table" else "room"
        self._drill = None
        self._load()

    def _on_drill_back(self):
        """面包屑返回：清下钻回到全局球房榜"""
        self._drill = None
        self._level_seg.blockSignals(True)
        self._level_seg.setCurrentItem("room")
        self._level_seg.blockSignals(False)
        self._level = "room"
        self._load()

    def _on_chart_row(self, idx: int):
        """图表行点击：球房全局榜=下钻该球房，其余=跳转明细"""
        if not (0 <= idx < len(self._rows)):
            return
        rec = self._rows[idx]
        if self._level == "room" and not self._drill:
            self._drill_into(rec)
        else:
            self._goto_detail(rec)

    def _on_ops_link(self, row: int, _col: int, action: str):
        """操作列链接点击分发（core.ops_link_delegate 回调）"""
        if not (0 <= row < len(self._rows)):
            return
        rec = self._rows[row]
        if action == "drill":
            self._drill_into(rec)
        elif action == "detail":
            self._goto_detail(rec)

    def _drill_into(self, rec: dict):
        room = str(rec.get("room_name") or "").strip()
        if not room:
            return
        self._drill = room
        self._level = "table"
        self._level_seg.blockSignals(True)
        self._level_seg.setCurrentItem("table")
        self._level_seg.blockSignals(False)
        self._load()

    def _goto_detail(self, rec: dict):
        """明细跳转：球桌级按桌号、球房级按球房名预筛选记录页
        （关键词搜索覆盖 table_no/room_name 字段）"""
        kw = (str(rec.get("table_no") or "").strip()
              if rec.get("table_no")
              else str(rec.get("room_name") or "").strip())
        if not kw:
            kw = str(rec.get("name") or "").strip()
        self.jump_to_records.emit(kw)

    # ---------- 导出 ----------

    def _on_export(self):
        """导出当前排行（含概览口径）为 xlsx（openpyxl，桌面端既有范式）"""
        if not self._rows:
            show_info_bar("当前排行无数据可导出", "warning",
                          title="导出", parent=self, duration=2500)
            return
        default_name = (f"售后排行_{'球桌' if (self._drill or self._level == 'table') else '球房'}_"
                        f"{date.today().isoformat()}.xlsx")
        path, _ = QFileDialog.getSaveFileName(
            self, "导出排行", default_name, "Excel 工作簿 (*.xlsx)")
        if not path:
            return
        try:
            n = self._write_rank_xlsx(path)
        except Exception as e:
            show_info_bar(f"导出失败：{type(e).__name__}: {e}", "error",
                          title="导出", parent=self, duration=4000)
            return
        show_info_bar(f"已导出 {n} 行排行 → {path}", "success",
                      title="导出", parent=self, duration=3000)

    def _write_rank_xlsx(self, path: str) -> int:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill

        wb = Workbook()
        ws = wb.active
        ws.title = "售后排行"
        headers = ("排名", "名称", "总量", "占比(%)", "未解决",
                   "我方问题", "主动发起", "最近发生")
        header_font = Font(bold=True)
        header_fill = PatternFill("solid", fgColor="D9F2F4")
        for c, h in enumerate(headers, 1):
            cell = ws.cell(row=1, column=c, value=h)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")
        for r, rec in enumerate(self._rows, 2):
            values = (
                int(rec.get("rank") or 0), str(rec.get("name") or ""),
                int(rec.get("total") or 0), int(rec.get("share") or 0),
                int(rec.get("unresolved") or 0),
                int(rec.get("our_problem") or 0),
                int(rec.get("initiative") or 0),
                str(rec.get("last_occurred") or ""))
            for c, v in enumerate(values, 1):
                ws.cell(row=r, column=c, value=v)
        s = self._summary
        last = len(self._rows) + 2
        summary_note = (f"口径：共 {int(s.get('total') or 0)} 条记录 · "
                        f"覆盖球房 {int(s.get('rooms') or 0)} · "
                        f"覆盖球桌 {int(s.get('tables') or 0)} · "
                        f"未解决 {int(s.get('unresolved') or 0)}")
        ws.cell(row=last, column=1, value=summary_note).font = header_font
        for c, w in enumerate((8, 32, 8, 8, 8, 10, 10, 12), 1):
            ws.column_dimensions[ws.cell(row=1, column=c).column_letter].width = w
        ws.freeze_panes = "A2"
        wb.save(path)
        return len(self._rows)
